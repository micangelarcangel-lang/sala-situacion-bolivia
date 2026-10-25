# -*- coding: utf-8 -*-
"""Base de datos descargable de la Sala de Situación: CSV, Excel, JSON y GeoJSON con diccionario y metadatos.
Se acumula entre ejecuciones (historico_eventos.csv) para que la base crezca con el tiempo."""
import os, csv, json, datetime as dt
import reporte_pdf as rp

NIVELES = ["verde", "amarilla", "naranja", "roja"]
CAMPOS_EV = ["id", "clase", "fecha_hora_local", "ts_ms", "fuente", "tipo_amenaza", "nivel_num", "nivel_preliminar", "titulo", "detalle", "url",
             "fallecidos", "heridos", "lat", "lon", "lugar", "metodo_ubicacion", "departamento", "municipio"]
CAMPOS_HE = ["hecho_id", "fecha_hora_local", "ts_ms", "titulo", "tipo_amenaza", "nivel_num", "nivel_preliminar", "fallecidos_max", "heridos_max",
             "cifras_discrepantes", "n_notas", "fuentes", "lat", "lon", "lugar", "departamento", "municipio", "enlaces"]

DICCIONARIO = [
    ("id", "texto", "Identificador único del evento en el registro del sistema.", "Sistema"),
    ("clase", "texto", "Tipo de registro: noticia, usgs, gdacs, clima, registro (ingresado a mano) u otro.", "Derivado del id"),
    ("fecha_hora_local", "fecha-hora", "Fecha y hora de publicación o detección, hora de Bolivia (UTC-4).", "Fuente original"),
    ("ts_ms", "entero", "Marca de tiempo en milisegundos desde 1970-01-01 UTC.", "Fuente original"),
    ("fuente", "texto", "Medio o servicio que reporta (p. ej. USGS, Open-Meteo, nombre del medio de prensa, Registro abierto).", "Fuente original"),
    ("tipo_amenaza", "texto", "Clase de amenaza (art. 36 Ley 602): geológica, meteorológica, hidrológica, climatológica, biológica, tecnológica, antropogénica u otra. Clasificación automática por palabras clave.", "Regla del sistema"),
    ("nivel_num", "0-3", "Nivel preliminar CCES: 0 verde, 1 amarilla, 2 naranja, 3 roja. NO es una alerta oficial de la Ley 602.", "Regla del sistema"),
    ("nivel_preliminar", "texto", "Nombre del nivel preliminar.", "Regla del sistema"),
    ("titulo", "texto", "Titular de la nota o descripción del evento.", "Fuente original"),
    ("detalle", "texto", "Información adicional (valores del modelo, descripción del registro).", "Fuente original"),
    ("url", "texto", "Enlace a la fuente. Los enlaces de Google Noticias redirigen al medio original.", "Fuente original"),
    ("fallecidos", "entero", "Fallecidos que menciona la fuente (cifra de prensa, sin validar).", "Texto de la fuente"),
    ("heridos", "entero", "Heridos que menciona la fuente (cifra de prensa, sin validar).", "Texto de la fuente"),
    ("lat", "decimal", "Latitud (WGS84). Puede ser aproximada: ver metodo_ubicacion.", "Fuente o gazetteer"),
    ("lon", "decimal", "Longitud (WGS84).", "Fuente o gazetteer"),
    ("lugar", "texto", "Nombre del lugar determinado.", "Fuente o gazetteer"),
    ("metodo_ubicacion", "texto", "Cómo se obtuvo la ubicación (coordenada de la fuente, municipio o departamento nombrado en el titular, sin ubicar).", "Sistema"),
    ("departamento", "texto", "Departamento según los límites geoBoundaries (se calcula a partir de lat/lon).", "geoBoundaries"),
    ("municipio", "texto", "Municipio según los límites geoBoundaries (se calcula a partir de lat/lon).", "geoBoundaries"),
    ("hecho_id", "texto", "(hechos) id del evento representativo del hecho.", "Sistema"),
    ("fallecidos_max / heridos_max", "entero", "(hechos) Máximo reportado entre las notas del mismo hecho. No se suman porque las notas cuentan a las mismas personas.", "Regla del sistema"),
    ("cifras_discrepantes", "sí/no", "(hechos) Las fuentes dan cifras distintas de fallecidos o heridos.", "Regla del sistema"),
    ("n_notas / fuentes / enlaces", "entero/texto", "(hechos) Notas unidas en el hecho, medios que las publicaron y sus enlaces.", "Sistema"),
]


def _hora(ms):
    return dt.datetime.fromtimestamp(ms / 1000, dt.timezone(dt.timedelta(hours=-4))).strftime("%Y-%m-%d %H:%M")


def _clase(i):
    p = str(i).split(":")[0].lower()
    return p if p in ("noticia", "usgs", "gdacs", "clima", "registro") else ("registro" if p == "manual" else "otro")


def _fila(e, geo):
    dep, mun = rp.ubicar(e, geo) if e.get("lat") is not None else (None, None)
    n = int(e.get("nivel") or 0)
    return {"id": e["id"], "clase": _clase(e["id"]), "fecha_hora_local": _hora(e["ts"]), "ts_ms": int(e["ts"]), "fuente": e.get("fuente", ""),
            "tipo_amenaza": e.get("tipo", ""), "nivel_num": n, "nivel_preliminar": NIVELES[min(max(n, 0), 3)], "titulo": e.get("titulo", ""),
            "detalle": e.get("detalle", ""), "url": e.get("url", ""), "fallecidos": int(e.get("fall") or 0), "heridos": int(e.get("her") or 0),
            "lat": "" if e.get("lat") is None else round(float(e["lat"]), 5), "lon": "" if e.get("lon") is None else round(float(e["lon"]), 5),
            "lugar": e.get("lugar", ""), "metodo_ubicacion": e.get("metodo", ""), "departamento": dep or "", "municipio": mun or ""}


def _a_evento(f):
    return {"id": f["id"], "ts": int(f["ts_ms"]), "fuente": f["fuente"], "tipo": f["tipo_amenaza"], "nivel": int(f["nivel_num"] or 0), "titulo": f["titulo"],
            "detalle": f["detalle"], "url": f["url"], "fall": int(f["fallecidos"] or 0), "her": int(f["heridos"] or 0),
            "lat": float(f["lat"]) if f["lat"] != "" else None, "lon": float(f["lon"]) if f["lon"] != "" else None, "lugar": f["lugar"], "metodo": f["metodo_ubicacion"]}


def _escribir_csv(ruta, campos, filas):
    with open(ruta, "w", encoding="utf-8-sig", newline="") as fh:  # utf-8-sig: Excel abre bien las tildes
        w = csv.DictWriter(fh, fieldnames=campos, extrasaction="ignore")
        w.writeheader()
        w.writerows(filas)


def exportar(carpeta, ahora_ms, eventos, geo, base_url=""):
    """eventos: lista de eventos ya filtrados (solo Bolivia). Devuelve resumen."""
    os.makedirs(carpeta, exist_ok=True)
    ruta_h = os.path.join(carpeta, "historico_eventos.csv")
    hist = {}
    if os.path.exists(ruta_h):
        for f in csv.DictReader(open(ruta_h, encoding="utf-8-sig", newline="")):
            hist[f["id"]] = f
    for e in eventos:
        hist[e["id"]] = _fila(e, geo)
    filas = sorted(hist.values(), key=lambda f: -int(f["ts_ms"]))
    _escribir_csv(ruta_h, CAMPOS_EV, filas)

    # hechos agrupados (último año)
    lim = ahora_ms - 365 * 86400000
    evs = [_a_evento(f) for f in filas if int(f["ts_ms"]) >= lim]
    hechos = []
    for r in rp.agrupar_hechos(evs):
        dep, mun = rp.ubicar(r, geo) if r.get("lat") is not None else (None, None)
        n = int(r["nivel"])
        hechos.append({"hecho_id": r["id"], "fecha_hora_local": _hora(r["ts"]), "ts_ms": r["ts"], "titulo": r["titulo"], "tipo_amenaza": r["tipo"], "nivel_num": n,
                       "nivel_preliminar": NIVELES[min(max(n, 0), 3)], "fallecidos_max": r.get("fall") or 0, "heridos_max": r.get("her") or 0,
                       "cifras_discrepantes": "sí" if r.get("_contradice") else "no", "n_notas": len(r.get("_ids", [r["id"]])),
                       "fuentes": "; ".join(r.get("_fuentes", [r["fuente"]])), "lat": "" if r.get("lat") is None else round(r["lat"], 5),
                       "lon": "" if r.get("lon") is None else round(r["lon"], 5), "lugar": r.get("lugar", ""), "departamento": dep or "", "municipio": mun or "",
                       "enlaces": " | ".join(u for _, u in r.get("_enlaces", []))})
    hechos.sort(key=lambda h: -h["ts_ms"])
    _escribir_csv(os.path.join(carpeta, "hechos_agrupados.csv"), CAMPOS_HE, hechos)
    json.dump(hechos, open(os.path.join(carpeta, "hechos_agrupados.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    gj = {"type": "FeatureCollection", "features": [
        {"type": "Feature", "geometry": {"type": "Point", "coordinates": [h["lon"], h["lat"]]},
         "properties": {k: v for k, v in h.items() if k not in ("lat", "lon")}} for h in hechos if h["lat"] != ""]}
    json.dump(gj, open(os.path.join(carpeta, "hechos.geojson"), "w", encoding="utf-8"), ensure_ascii=False)
    _escribir_csv(os.path.join(carpeta, "diccionario_datos.csv"), ["campo", "tipo", "descripcion", "origen"],
                  [dict(zip(["campo", "tipo", "descripcion", "origen"], d)) for d in DICCIONARIO])

    meta = {"generado_local": _hora(ahora_ms), "zona_horaria": "America/La_Paz (UTC-4)", "registros_eventos": len(filas), "registros_hechos": len(hechos),
            "alcance": "Solo hechos de Bolivia. Datos preliminares de fuentes abiertas, sin validación oficial.",
            "fuentes": ["USGS (sismos)", "GDACS (desastres)", "Open-Meteo (pronóstico, CC BY 4.0)", "Google Noticias (titulares y enlaces; derechos de cada medio)",
                        "Registro abierto de la Sala de Situación (sin verificar)", "Límites: geoBoundaries"],
            "advertencias": ["El nivel de color es preliminar (criterio CCES) y no corresponde a una declaratoria de alerta de la Ley 602.",
                             "Fallecidos y heridos son cifras de prensa, sin validar. En un hecho se toma el máximo reportado, no la suma.",
                             "La ubicación puede ser aproximada (ver metodo_ubicacion)."],
            "cita_sugerida": "Sala de Situación Bolivia – CCES. Base de eventos de amenazas (datos preliminares de fuentes abiertas). Consulta: " + _hora(ahora_ms) + "."}
    json.dump(meta, open(os.path.join(carpeta, "metadatos.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    xlsx_ok = True
    try:
        from openpyxl import Workbook
        wb = Workbook()
        for nombre, campos, datos in (("Hechos", CAMPOS_HE, hechos), ("Eventos", CAMPOS_EV, filas)):
            ws = wb.active if nombre == "Hechos" else wb.create_sheet(nombre)
            ws.title = nombre
            ws.append(campos)
            for d in datos:
                ws.append([d.get(c, "") for c in campos])
            ws.freeze_panes = "A2"
        ws = wb.create_sheet("Diccionario"); ws.append(["campo", "tipo", "descripcion", "origen"])
        for d in DICCIONARIO:
            ws.append(list(d))
        ws = wb.create_sheet("Metadatos")
        for k, v in meta.items():
            ws.append([k, "; ".join(v) if isinstance(v, list) else v])
        wb.save(os.path.join(carpeta, "datos_bolivia.xlsx"))
    except Exception as e:
        xlsx_ok = False
        print("No se generó el Excel:", type(e).__name__)

    arch = [("hechos_agrupados.csv", "Hechos (una fila por hecho) — CSV"), ("datos_bolivia.xlsx", "Excel con hechos, eventos, diccionario y metadatos"),
            ("historico_eventos.csv", "Todos los eventos acumulados — CSV"), ("hechos.geojson", "Hechos con ubicación — GeoJSON (QGIS, ArcGIS)"),
            ("hechos_agrupados.json", "Hechos — JSON"), ("diccionario_datos.csv", "Diccionario de datos"), ("metadatos.json", "Metadatos y cita sugerida")]
    if not xlsx_ok:
        arch = [a for a in arch if not a[0].endswith(".xlsx")]
    li = "".join(f'<li><a href="{f}">{t}</a> <small>({os.path.getsize(os.path.join(carpeta, f)) // 1024} KB)</small></li>' for f, t in arch)
    open(os.path.join(carpeta, "index.html"), "w", encoding="utf-8").write(
        '<!doctype html><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Datos abiertos</title>'
        '<style>body{font:15px/1.6 system-ui;max-width:760px;margin:24px auto;padding:0 16px}a{color:#0b5cab}small{color:#666}</style>'
        '<h1>Datos — Sala de Situación Bolivia</h1><p><a href="../">← Volver al tablero</a></p>'
        f'<p>Actualizado: <b>{meta["generado_local"]}</b> (hora de Bolivia). {len(filas)} eventos, {len(hechos)} hechos.</p><ul>{li}</ul>'
        '<h3>Antes de usar los datos</h3><ul>' + "".join(f"<li>{a}</li>" for a in meta["advertencias"]) + "</ul>"
        f'<p><b>Cómo citar:</b> {meta["cita_sugerida"]}</p>')
    return {"eventos": len(filas), "hechos": len(hechos), "xlsx": xlsx_ok}
